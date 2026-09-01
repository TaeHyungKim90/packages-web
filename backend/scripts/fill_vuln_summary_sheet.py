"""Fill sheet '5점이상 취약점' from pypi/npm/nuget vulnerability sheets."""

from __future__ import annotations

import asyncio
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

# Allow running as `uv run python scripts/...` from backend/
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import ECOSYSTEM_MAP, settings  # noqa: E402
from app.services.nexus import _auth, _display_name  # noqa: E402
from app.services.osv import (  # noqa: E402
    _is_newer,
    _version_key,
    is_package_fixed_version,
)

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "패키지 취약점 문제_v4_2026-08-27.xlsx"
SHEET4 = "5점이상 취약점"
SOURCE_SHEETS = ("pypi", "npm", "nuget")
MIN_SCORE = 5.0

NON_VERSION_FIXED = re.compile(
    r"(해당\s*없|해결\s*불|해결불가|오탐|없음|n/?a|none|unfixed|no\s*fix)",
    re.I,
)


@dataclass
class PackageAgg:
    name: str
    vulnerable: set[str] = field(default_factory=set)
    fixed_raw: set[str] = field(default_factory=set)
    max_score: float = 0.0


def _split_versions(cell: str | None) -> list[str]:
    if not cell:
        return []
    text = str(cell).strip()
    if not text:
        return []
    parts = re.split(r"[,;/|]\s*", text)
    return [p.strip() for p in parts if p.strip()]


def _major(version: str) -> str:
    m = re.match(r"^v?(\d+)", version.strip(), re.I)
    return m.group(1) if m else ""


def _is_stable(version: str) -> bool:
    text = version.strip()
    if not text or not is_package_fixed_version(text):
        return False
    if "-" in text or "+" in text:
        return False
    lower = text.lower()
    if re.search(r"(a|b|rc|dev|alpha|beta|pre|preview)\d*$", lower):
        return False
    return True


def _pick_fixed_representatives(fixed_values: set[str]) -> list[str]:
    """Per major line keep the highest fix; non-version notes kept as-is once."""
    notes: list[str] = []
    by_major: dict[str, list[str]] = defaultdict(list)
    for raw in fixed_values:
        text = raw.strip()
        if not text:
            continue
        if NON_VERSION_FIXED.search(text) or not is_package_fixed_version(text):
            if text not in notes:
                notes.append(text)
            continue
        for ver in _split_versions(text):
            if NON_VERSION_FIXED.search(ver) or not is_package_fixed_version(ver):
                if ver not in notes:
                    notes.append(ver)
                continue
            maj = _major(ver) or "_"
            by_major[maj].append(ver)
    picked: list[str] = []
    for maj in sorted(by_major, key=lambda m: (m == "_", int(m) if m.isdigit() else 999)):
        versions = by_major[maj]
        picked.append(max(versions, key=_version_key))
    # Prefer concrete versions first, then notes
    return sorted(picked, key=_version_key) + notes


def _sort_versions(versions: set[str]) -> list[str]:
    def key(v: str):
        try:
            return (0, _version_key(v))
        except Exception:
            return (1, v.lower())

    return sorted(versions, key=key)


def _header_map(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value is None:
            continue
        headers[str(cell.value).strip()] = cell.column
    return headers


def load_aggregates(wb) -> dict[str, dict[str, PackageAgg]]:
    out: dict[str, dict[str, PackageAgg]] = {}
    for eco in SOURCE_SHEETS:
        ws = wb[eco]
        headers = _header_map(ws)
        need = ("위험도", "이름", "버전", "해결 버전")
        for h in need:
            if h not in headers:
                raise RuntimeError(f"{eco}: missing column {h}: {headers}")
        c_score = headers["위험도"]
        c_name = headers["이름"]
        c_ver = headers["버전"]
        c_fix = headers["해결 버전"]
        pkgs: dict[str, PackageAgg] = {}
        for row in range(2, ws.max_row + 1):
            name = ws.cell(row, c_name).value
            if not name:
                continue
            name = str(name).strip()
            score_raw = ws.cell(row, c_score).value
            try:
                score = float(score_raw) if score_raw is not None else 0.0
            except (TypeError, ValueError):
                score = 0.0
            if score < MIN_SCORE:
                continue
            agg = pkgs.get(name)
            if agg is None:
                agg = PackageAgg(name=name)
                pkgs[name] = agg
            agg.max_score = max(agg.max_score, score)
            for v in _split_versions(ws.cell(row, c_ver).value):
                agg.vulnerable.add(v)
            fix_cell = ws.cell(row, c_fix).value
            if fix_cell is not None and str(fix_cell).strip():
                agg.fixed_raw.add(str(fix_cell).strip())
        out[eco] = pkgs
    return out


def _name_matches(query: str, found: str, eco: str) -> bool:
    q = query.strip()
    f = found.strip()
    if q.lower() == f.lower():
        return True
    if eco == "pypi":
        def norm(s: str) -> str:
            return re.sub(r"[-_.]+", "-", s).lower()

        return norm(q) == norm(f)
    if eco == "nuget":
        return q.lower() == f.lower()
    return q == f


async def _nexus_versions_for_package(
    client: httpx.AsyncClient,
    *,
    eco: str,
    package_name: str,
) -> list[str]:
    cfg = ECOSYSTEM_MAP[eco]
    params: dict[str, str] = {
        "repository": cfg.hosted_repo,
        "format": eco if eco != "pypi" else "pypi",
        "q": package_name,
    }
    # Prefer exact name filters when Nexus supports them
    if eco == "npm":
        if package_name.startswith("@") and "/" in package_name:
            scope, name = package_name.split("/", 1)
            params["group"] = scope.lstrip("@")
            params["name"] = name
        else:
            params["name"] = package_name
    elif eco == "nuget":
        params["name"] = package_name
    elif eco == "pypi":
        params["name"] = package_name

    versions: set[str] = set()
    token: str | None = None
    for _ in range(20):
        page = dict(params)
        if token:
            page["continuationToken"] = token
        url = f"{settings.nexus_base_url.rstrip('/')}/service/rest/v1/search"
        response = await client.get(url, params=page, auth=_auth(), timeout=60.0)
        response.raise_for_status()
        data = response.json()
        for raw in data.get("items") or []:
            disp = _display_name(raw)
            ver = raw.get("version")
            if not disp or not ver:
                continue
            if not _name_matches(package_name, disp, eco):
                continue
            versions.add(str(ver))
        token = data.get("continuationToken")
        if not token:
            break
    return _sort_versions(versions)


def _meets_any_fixed(
    candidate: str,
    fixed_list: list[str],
    *,
    vulnerable: set[str],
) -> bool:
    if candidate in vulnerable:
        return False
    concrete = [
        f
        for f in fixed_list
        if is_package_fixed_version(f) and not NON_VERSION_FIXED.search(f)
    ]
    if not concrete:
        return False
    maj = _major(candidate)
    same_major = [f for f in concrete if maj and _major(f) == maj]
    if same_major:
        need = max(same_major, key=_version_key)
        return candidate == need or _is_newer(candidate, need)
    # No fix published on this major: only accept if newer than every known fix.
    highest = max(concrete, key=_version_key)
    return candidate == highest or _is_newer(candidate, highest)


async def fill_available(
    aggregates: dict[str, dict[str, PackageAgg]],
) -> dict[str, dict[str, str]]:
    """eco -> name -> comma-joined available versions."""
    result: dict[str, dict[str, str]] = {eco: {} for eco in SOURCE_SHEETS}
    verify = bool(settings.nexus_verify_ssl)
    async with httpx.AsyncClient(verify=verify) as client:
        for eco, pkgs in aggregates.items():
            print(f"[nexus] {eco}: {len(pkgs)} packages")
            for i, (name, agg) in enumerate(sorted(pkgs.items(), key=lambda x: x[0].lower()), 1):
                fixed = _pick_fixed_representatives(agg.fixed_raw)
                try:
                    hosted = await _nexus_versions_for_package(
                        client, eco=eco, package_name=name
                    )
                except Exception as exc:
                    print(f"  ! {name}: {exc}")
                    result[eco][name] = ""
                    continue
                safe = [
                    v
                    for v in hosted
                    if _is_stable(v)
                    and _meets_any_fixed(v, fixed, vulnerable=agg.vulnerable)
                ]
                result[eco][name] = ", ".join(safe)
                if i % 10 == 0 or i == len(pkgs):
                    print(f"  … {i}/{len(pkgs)}")
    return result


def write_sheet4(
    wb,
    aggregates: dict[str, dict[str, PackageAgg]],
    available: dict[str, dict[str, str]],
) -> None:
    # Drop any existing summary sheets (name may have trailing space).
    for name in list(wb.sheetnames):
        if name.strip() == SHEET4.strip():
            wb.remove(wb[name])
    ws = wb.create_sheet(SHEET4)

    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_font = Font(bold=True)
    title_font = Font(bold=True, size=12)
    headers = ["이름", "버전", "해결된 최소버전", "현재있는버전"]

    row = 1
    for eco in SOURCE_SHEETS:
        ws.cell(row, 1, eco).font = title_font
        row += 1
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row, col, h)
            cell.font = header_font
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        row += 1

        pkgs = aggregates.get(eco) or {}
        for name in sorted(pkgs, key=str.lower):
            agg = pkgs[name]
            fixed = _pick_fixed_representatives(agg.fixed_raw)
            values = [
                name,
                ", ".join(_sort_versions(agg.vulnerable)),
                ", ".join(fixed),
                available.get(eco, {}).get(name, ""),
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row, col, val)
                cell.border = thin
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
        row += 1  # blank between ecosystems

    widths = [28, 40, 28, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


async def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"missing workbook: {XLSX}")
    print(f"loading {XLSX}")
    wb = openpyxl.load_workbook(XLSX)
    aggregates = load_aggregates(wb)
    for eco, pkgs in aggregates.items():
        print(f"{eco}: {len(pkgs)} packages (>= {MIN_SCORE})")
    available = await fill_available(aggregates)
    write_sheet4(wb, aggregates, available)
    wb.save(XLSX)
    print(f"saved {XLSX}")

    # Spot checks
    spots = (
        ("npm", "axios"),
        ("pypi", "cryptography"),
        ("nuget", "Microsoft.IdentityModel.Tokens.Saml"),
    )
    for eco, name in spots:
        agg = aggregates.get(eco, {}).get(name)
        if not agg:
            print(f"spot {eco}/{name}: missing")
            continue
        print(
            f"spot {eco}/{name}: vuln={_sort_versions(agg.vulnerable)} "
            f"fix={_pick_fixed_representatives(agg.fixed_raw)} "
            f"avail={available.get(eco, {}).get(name)!r}"
        )


if __name__ == "__main__":
    asyncio.run(main())
